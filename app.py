
from __future__ import annotations

import io
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Previsão Financeira", page_icon="💰", layout="wide")

# ============================================================
# Utilidades
# ============================================================

def money(v: float) -> str:
    return f"R$ {v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

def parse_money(v) -> float:
    if v is None:
        return 0.0
    s = str(v).strip().replace("R$", "").replace(" ", "")
    if not s:
        return 0.0
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except Exception:
        return 0.0

def parse_br_date(v):
    if pd.isna(v) or str(v).strip() == "":
        return pd.NaT
    return pd.to_datetime(str(v).strip(), dayfirst=True, errors="coerce").date()

# ============================================================
# Leitor do cubo bruto
# Estrutura observada:
# carteira;codigo;data;tipo;valor fornecedor;valor cliente
# As linhas seguintes podem deixar carteira/código em branco.
# ============================================================

def ler_cubo_bruto(uploaded_file) -> pd.DataFrame:
    """
    Lê automaticamente os dois layouts de saída encontrados no Tecnicon:

    1) Cubo financeiro atual:
       Carteira;Cód. Carteira;Data Vencimento;Tipo Carteira;
       Duplicata;Cliente/Fornecedor;Valor Fornecedor;Valor Cliente

    2) Relatório "DUPLICATAS DE FORNECEDORES A VENCER":
       CODIGO;FL;FORNECEDOR;DUPLICATA;PARCELA;VALOR;IRRF NF;
       VLR EM MOEDA;CODIGO;CARTEIRA;FL

    O segundo layout usa blocos "VENCIMENTO: dd/mm/aaaa;". A data é
    herdada pelas linhas seguintes até aparecer um novo vencimento.
    """
    dados = uploaded_file.getvalue()
    try:
        raw = dados.decode("utf-8-sig")
    except UnicodeDecodeError:
        raw = dados.decode("cp1252", errors="replace")

    # ------------------------------------------------------------
    # Layout 2: DUPLICATAS DE FORNECEDORES A VENCER
    # ------------------------------------------------------------
    if "DUPLICATAS DE FORNECEDORES A VENCER" in raw.upper():
        registros = []
        data_vencimento = pd.NaT

        for linha in raw.splitlines():
            linha = linha.strip()
            if not linha:
                continue

            # Ex.: VENCIMENTO: 09/09/2026;
            if linha.upper().startswith("VENCIMENTO:"):
                valor_data = linha.split(":", 1)[1].strip().rstrip(";").strip()
                parsed = parse_br_date(valor_data)
                if not pd.isna(parsed):
                    data_vencimento = parsed
                continue

            partes = [p.strip() for p in linha.split(";")]
            if not partes:
                continue

            primeira = partes[0].upper()
            # Ignora títulos, cabeçalhos, separadores e totais.
            if (
                primeira.startswith("SOBERANA ALIMENTOS")
                or primeira.startswith("DUPLICATAS DE FORNECEDORES")
                or primeira.startswith("CODIGO")
                or primeira.startswith("_")
                or primeira.startswith("TOTAL VENCIMENTO")
            ):
                continue

            while len(partes) < 11:
                partes.append("")

            # Uma linha de lançamento possui fornecedor + duplicata + valor.
            # O Tecnicon pode omitir o campo "VLR EM MOEDA" quando ele está
            # vazio. Por isso, as posições finais são lidas pelo conteúdo: os
            # três últimos campos não vazios são CODIGO, CARTEIRA e FL.
            fornecedor_nome = partes[2]
            duplicata = partes[3]
            valor = parse_money(partes[5])

            finais = [p for p in partes[6:] if p.strip()]
            if len(finais) < 3:
                continue
            carteira_codigo, carteira_nome, _fl = finais[-3:]

            if pd.isna(data_vencimento):
                continue
            if not fornecedor_nome or not duplicata:
                continue
            if valor == 0:
                continue

            registros.append({
                "Carteira": carteira_codigo,
                "Código": carteira_codigo,
                "Data": data_vencimento,
                "Tipo": "FORNECEDORES",
                "Fornecedor": fornecedor_nome,
                "Duplicata": duplicata,
                "Valor Fornecedor": valor,
                "Valor Cliente": 0.0,
                "Carteira Nome": carteira_nome,
            })

        return pd.DataFrame(registros)

    # ------------------------------------------------------------
    # Layout 1: Cubo financeiro atual / antigo
    # ------------------------------------------------------------
    registros = []
    carteira = ""
    codigo = ""
    tipo = ""
    data_atual = pd.NaT
    duplicata = ""

    for linha in raw.splitlines():
        linha = linha.strip()
        if not linha:
            continue

        partes = [p.strip() for p in linha.split(";")]
        n_campos_original = len(partes)

        while len(partes) > 0 and partes[-1] == "":
            partes.pop()

        if partes and "CarteiraCód." in partes[0]:
            continue

        if n_campos_original >= 8:
            while len(partes) < 8:
                partes.append("")

            novo_bloco = bool(partes[0])
            if novo_bloco:
                carteira = partes[0]
                codigo = partes[1]
                tipo = partes[3].strip() if partes[3].strip() else ""

            if partes[2]:
                parsed = parse_br_date(partes[2])
                if not pd.isna(parsed):
                    data_atual = parsed

            if pd.isna(data_atual):
                continue

            duplicata = partes[4]
            fornecedor_nome = partes[5]
            fornecedor = parse_money(partes[6])
            cliente = parse_money(partes[7])

        else:
            while len(partes) < 6:
                partes.append("")

            if partes[0]:
                carteira = partes[0]
                codigo = partes[1]
                tipo = partes[3].strip() if partes[3].strip() else ""

            if partes[2]:
                parsed = parse_br_date(partes[2])
                if not pd.isna(parsed):
                    data_atual = parsed

            if pd.isna(data_atual):
                continue

            fornecedor_nome = ""
            duplicata = ""
            fornecedor = parse_money(partes[4])
            cliente = parse_money(partes[5])

        registros.append({
            "Carteira": carteira,
            "Código": codigo,
            "Data": data_atual,
            "Tipo": tipo,
            "Fornecedor": fornecedor_nome,
            "Duplicata": duplicata,
            "Valor Fornecedor": fornecedor,
            "Valor Cliente": cliente,
        })

    df_result = pd.DataFrame(registros)

    if not df_result.empty:
        df_result["Tipo"] = df_result["Tipo"].fillna("").astype(str).str.strip()
        tipo_upper = df_result["Tipo"].str.upper()
        desconhecido = ~tipo_upper.str.startswith(("CLIENT", "FORNECED", "IMPOSTOS"))
        fornecedor_pos = pd.to_numeric(df_result["Valor Fornecedor"], errors="coerce").fillna(0) > 0
        cliente_pos = pd.to_numeric(df_result["Valor Cliente"], errors="coerce").fillna(0) > 0
        df_result.loc[desconhecido & fornecedor_pos & ~cliente_pos, "Tipo"] = "FORNECEDORES"
        df_result.loc[desconhecido & cliente_pos & ~fornecedor_pos, "Tipo"] = "CLIENTES"

    return df_result


def eh_cliente(tipo):
    return str(tipo).strip().upper().startswith("CLIENT")


def eh_fornecedor(tipo):
    return str(tipo).strip().upper().startswith("FORNECED")


def eh_imposto(tipo):
    return "IMPOSTOS" in str(tipo).strip().upper()


# ============================================================
# Dias bancários
# D+0 = a própria data se compensável.
# D+1 = primeiro dia útil/compensável seguinte.
# Feriados cadastrados manualmente são ignorados.
# ============================================================

def proxima_compensacao(data_base, dplus, sem_comp):
    atual = data_base
    passos = 0

    if dplus == 0:
        while atual.weekday() >= 5 or atual in sem_comp:
            atual += timedelta(days=1)
        return atual

    while passos < dplus:
        atual += timedelta(days=1)
        if atual.weekday() < 5 and atual not in sem_comp:
            passos += 1

    while atual.weekday() >= 5 or atual in sem_comp:
        atual += timedelta(days=1)

    return atual

# ============================================================
# Excel de saída, no estilo da planilha antiga
# ============================================================

def gerar_excel(df, saldo_inicial, data_inicio, config, sem_comp, data_fim=None):
    """
    Usa a planilha fornecida pelo usuário como MODELO.

    Fluxo:
      CSV -> Detalhamento/Base -> fórmulas da planilha -> Previsão Financeira

    O saldo inicial é manual. Recebimentos usam Data Entrada (D+).
    Pagamentos usam Data Saída/data original.
    """
    from copy import copy
    from openpyxl import load_workbook

    modelo = Path(__file__).with_name("Previsao_Financeira_Modelo.xlsx")
    if not modelo.exists():
        raise FileNotFoundError(
            "O arquivo Previsao_Financeira_Modelo.xlsx não foi encontrado "
            "junto ao app.py."
        )

    # ------------------------------------------------------------
    # Preparar dados
    # ------------------------------------------------------------
    work = df.copy()

    # Normaliza tipos para evitar problemas com datas/códigos.
    work["Código"] = work["Código"].astype(str).str.strip()
    work["Carteira"] = work["Carteira"].astype(str).str.strip()
    if "Fornecedor" not in work.columns:
        work["Fornecedor"] = ""
    work["Fornecedor"] = work["Fornecedor"].fillna("").astype(str).str.strip()
    if "Duplicata" not in work.columns:
        work["Duplicata"] = ""
    work["Duplicata"] = work["Duplicata"].fillna("").astype(str).str.strip()

    # Configuração D+ por código de carteira.
    cfg = {}
    for _, r in config.iterrows():
        carteira = str(r.get("Carteira", "")).strip()
        if carteira:
            try:
                cfg[carteira] = int(r.get("D+", 1))
            except Exception:
                cfg[carteira] = 1

    work["D+"] = work["Código"].map(cfg).fillna(1).astype(int)

    # Somente CLIENTES recebem D+.
    work["Data Entrada"] = work.apply(
        lambda r: proxima_compensacao(
            r["Data"], int(r["D+"]), sem_comp
        ) if eh_cliente(r["Tipo"]) else pd.NaT,
        axis=1,
    )

    # Contas a pagar: data exata do lançamento.
    work["Data Saída"] = work["Data"]

    # O leitor do CSV trabalha inicialmente com objetos `date`.
    # Antes de usar o acessador `.dt`, convertemos explicitamente as
    # colunas de data para datetime64. Isso evita:
    # "Can only use .dt accessor with datetimelike values".
    work["Data"] = pd.to_datetime(work["Data"], errors="coerce")
    work["Data Entrada"] = pd.to_datetime(work["Data Entrada"], errors="coerce")
    work["Data Saída"] = pd.to_datetime(work["Data Saída"], errors="coerce")

    # ------------------------------------------------------------
    # Abrir o modelo do usuário
    # ------------------------------------------------------------
    wb = load_workbook(modelo)
    # Excel deve recalcular as fórmulas ao abrir o arquivo gerado.
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass

    # A aba criada pelo usuário para o cálculo.
    if "Planilha1" in wb.sheetnames:
        ws = wb["Planilha1"]
    elif "Previsão Financeira" in wb.sheetnames:
        ws = wb["Previsão Financeira"]
    else:
        raise ValueError(
            "O modelo precisa ter a aba 'Planilha1' ou 'Previsão Financeira'."
        )

    # ------------------------------------------------------------
    # Remover bases antigas do modelo e recriar o Detalhamento
    # ------------------------------------------------------------
    if "Detalhamento" in wb.sheetnames:
        old_det = wb["Detalhamento"]
        wb.remove(old_det)

    det = wb.create_sheet("Detalhamento")

    headers = [
        "Data bruto",
        "Carteira",
        "Código",
        "Tipo",
        "D+",
        "Cliente",
        "Entrada banco",
        "Fornecedor/A pagar",
        "Saída banco",
        "Fornecedor",
        "Duplicata",
    ]
    det.append(headers)

    for cell in det[1]:
        # reaproveita estilo do cabeçalho se disponível
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")

    # Escreve a base completa.
    # As regras viajam dentro do Excel:
    #  - D+ (col. E) busca a carteira em Config_Carteiras (padrão D+1).
    #  - Entrada banco (col. G) aplica D+ só a CLIENTES, pulando fins de
    #    semana e as datas de Dias_Sem_Compensacao.
    #  Alterou a configuração no Excel -> tudo recalcula.
    FERIADOS = "Dias_Sem_Compensacao!$A$2:$A$500"
    for i, (_, r) in enumerate(work.iterrows(), start=2):
        f_dmais = (
            f'=IF(LEFT(D{i},6)="CLIENT",'
            f'IFERROR(VLOOKUP(C{i},Config_Carteiras!$A:$B,2,FALSE),1),"")'
        )
        f_entrada = (
            f'=IF(LEFT(D{i},6)="CLIENT",'
            f'IF(E{i}=0,WORKDAY(A{i}-1,1,{FERIADOS}),WORKDAY(A{i},E{i},{FERIADOS})),"")'
        )
        det.append([
            r["Data"],
            r["Carteira"],
            r["Código"],
            r["Tipo"],
            f_dmais,
            float(r.get("Valor Cliente", 0) or 0),
            f_entrada,
            float(r.get("Valor Fornecedor", 0) or 0),
            r["Data Saída"] if not pd.isna(r["Data Saída"]) else None,
            r.get("Fornecedor", ""),
            r.get("Duplicata", ""),
        ])

    for row in det.iter_rows(min_row=2):
        row[0].number_format = "dd/mm/yyyy"
        row[5].number_format = '#,##0.00'
        row[6].number_format = "dd/mm/yyyy"
        row[7].number_format = '#,##0.00'
        row[8].number_format = "dd/mm/yyyy"

    widths = [14, 32, 12, 24, 8, 18, 16, 20, 16, 48, 18]
    for i, width in enumerate(widths, start=1):
        det.column_dimensions[get_column_letter(i)].width = width

    det.freeze_panes = "A2"
    det.auto_filter.ref = f"A1:K{max(det.max_row, 2)}"

    # ------------------------------------------------------------
    # Período: exatamente o intervalo escolhido.
    # Máximo de 5 dias para manter o layout do Excel.
    # ------------------------------------------------------------
    segunda = data_inicio
    if data_fim is None:
        data_fim = data_inicio + timedelta(days=4)
    dias = [data_inicio + timedelta(days=i) for i in range((data_fim - data_inicio).days + 1)]

    # A planilha do usuário tem B4:F4 como datas.
    nomes_semana = ["SEGUNDA-FEIRA", "TERÇA-FEIRA", "QUARTA-FEIRA", "QUINTA-FEIRA",
                    "SEXTA-FEIRA", "SÁBADO", "DOMINGO"]
    for col_idx, dia in enumerate(dias, start=2):
        col = get_column_letter(col_idx)
        ws[f"{col}3"] = f"{dia.strftime('%d/%m')} - {nomes_semana[dia.weekday()].title()}"
        ws[f"{col}3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws[f"{col}4"] = dia
        ws[f"{col}4"].number_format = "dd/mm/yyyy"

    # A linha 4 continua contendo as datas reais para as fórmulas, mas fica
    # oculta para o usuário; a linha 3 mostra "dd/mm - Dia-da-semana".
    ws.row_dimensions[4].hidden = True

    # Limpa as colunas que ficarem fora do intervalo escolhido (até F).
    for c_idx in range(2, 7):
        if c_idx > len(dias) + 1:
            col = get_column_letter(c_idx)
            for rr in [3, 4, 8, 9, 10, 11, 13, 14, 15]:
                ws[f"{col}{rr}"] = None

    # Saldo inicial continua manual.
    ws["B6"] = float(saldo_inicial)
    ws["B6"].number_format = '#,##0.00'

    # ------------------------------------------------------------
    # Limpar valores antigos da área dinâmica da previsão.
    # Mantemos rótulos, estilos e estrutura da planilha.
    # ------------------------------------------------------------
    # Recebimentos: linhas 9-17.
    for r in range(9, 18):
        for c in range(2, 7):
            ws.cell(r, c).value = None

    # ------------------------------------------------------------
    # Recebimentos / cobranças bancárias
    # ------------------------------------------------------------
    # Linha 8 = total de todos os clientes pela DATA DE ENTRADA (D+).
    # Linhas 9-11 = detalhamento das principais cobranças bancárias.
    recebimento_linhas = {
        9: ("Cobrança Bradesc", "BRADESCO COBRANCA"),
        10: ("Cobrança Banrisul", "BANRISUL COBRANCA"),
        11: ("Depósito Clientes", "DEPOSITO (CLIENTE)"),
    }

    for row, (rotulo, carteira_busca) in recebimento_linhas.items():
        ws[f"A{row}"] = rotulo
        for c_idx in range(2, 7):
            col = get_column_letter(c_idx)
            ws[f"{col}{row}"] = (
                f'=SUMIFS(Detalhamento!$F:$F,'
                f'Detalhamento!$G:$G,{col}$4,'
                f'Detalhamento!$B:$B,"{carteira_busca}")'
            )
            ws[f"{col}{row}"].number_format = '#,##0.00'

    for c_idx in range(2, len(dias) + 2):
        col = get_column_letter(c_idx)
        ws[f"{col}8"] = (
            f'=SUMIF(Detalhamento!$G:$G,{col}$4,Detalhamento!$F:$F)'
        )
        ws[f"{col}8"].number_format = '#,##0.00'

    # "Outras carteiras": qualquer carteira que apareça no Cubo além das 3 já
    # detalhadas acima cai automaticamente aqui (total menos as 3 já listadas).
    # Não é preciso catalogar carteira nova em lugar nenhum do código.
    ws["A12"] = "Outras carteiras"
    for c_idx in range(2, len(dias) + 2):
        col = get_column_letter(c_idx)
        ws[f"{col}12"] = f"={col}8-{col}9-{col}10-{col}11"
        ws[f"{col}12"].number_format = '#,##0.00'

    # Pagamentos: reconstruímos os blocos dinamicamente para que TODOS
    # os fornecedores possam aparecer, mesmo quando houver mais lançamentos
    # do que as linhas existentes no modelo.
    #
    # Capturamos os estilos antes de apagar a área antiga.
    payment_style_rows = {
        "date": 19,
        "item": 20,
        "total": 25,
        "blank": 24,
    }
    payment_styles = {
        key: [copy(ws.cell(payment_style_rows[key], c)._style) for c in range(1, 7)]
        for key in payment_style_rows
    }
    payment_heights = {
        key: ws.row_dimensions[payment_style_rows[key]].height
        for key in payment_style_rows
    }

    # A partir da linha 19 só existem os blocos de contas a pagar no
    # modelo atual. Recriamos essa área sem deixar linhas antigas.
    if ws.max_row >= 19:
        ws.delete_rows(19, ws.max_row - 18)

    def aplicar_estilo_pagamento(row_num, tipo):
        estilos = payment_styles[tipo]
        for c in range(1, 7):
            ws.cell(row_num, c)._style = copy(estilos[c - 1])
        h = payment_heights.get(tipo)
        if h is not None:
            ws.row_dimensions[row_num].height = h

    # ------------------------------------------------------------
    # CONTAS A PAGAR
    # ------------------------------------------------------------
    # O novo Cubo traz o nome do fornecedor em cada linha.
    # A previsão lista cada fornecedor uma vez por dia; se houver mais de
    # um lançamento para o mesmo fornecedor no mesmo dia, os valores são
    # somados.
    current_row = 19

    for dia in dias:
        start_row = current_row
        aplicar_estilo_pagamento(start_row, "date")

        dia_ts = pd.Timestamp(dia).normalize()
        col = get_column_letter(dias.index(dia) + 2)
        ws[f"A{start_row}"] = f"={col}3"
        ws[f"A{start_row}"].number_format = "@"

        registros = work[
            work["Data Saída"].notna()
            & (work["Data Saída"].dt.normalize() == dia_ts)
            & (pd.to_numeric(work["Valor Fornecedor"], errors="coerce").fillna(0) != 0)
            & (work["Tipo"].map(eh_fornecedor) | work["Tipo"].map(eh_imposto))
        ].copy()

        # Mantemos cada lançamento individualmente. Isso permite que as
        # pequenas tabelas sejam ligadas diretamente ao Detalhamento.
        # Assim, se o usuário alterar o valor, duplicata ou nome na base,
        # a pequena tabela e o total são recalculados pelo Excel.
        if not registros.empty:
            registros["Fornecedor"] = registros["Fornecedor"].fillna("").astype(str).str.strip()
            registros["Duplicata"] = registros["Duplicata"].fillna("").astype(str).str.strip()
            registros["Nome Exibicao"] = registros.apply(
                lambda r: ("[IMPOSTO] " if eh_imposto(r["Tipo"]) else "") + (r["Fornecedor"] or "(sem nome)"),
                axis=1,
            )
            registros = registros.sort_values(by=["Nome Exibicao", "Duplicata"], kind="stable")

        item_start = start_row + 1

        for idx, (orig_idx, r) in enumerate(registros.iterrows()):
            row = item_start + idx
            aplicar_estilo_pagamento(row, "item")

            # O DataFrame foi escrito no Detalhamento a partir da linha 2,
            # preservando a ordem original. orig_idx é o índice original
            # do DataFrame e, portanto, aponta para a linha correspondente
            # na base do Excel.
            source_row = int(work.index.get_loc(orig_idx)) + 2

            # Tabela: Duplicata | Nome | Valor
            # As três células são fórmulas ligadas ao Detalhamento.
            ws[f"A{row}"] = f'=Detalhamento!K{source_row}'
            ws[f"B{row}"] = (
                f'=IF(Detalhamento!D{source_row}="IMPOSTOS E CONTRIBUICOES",'
                f'"[IMPOSTO] "&Detalhamento!J{source_row},Detalhamento!J{source_row})'
            )
            ws[f"C{row}"] = f'=Detalhamento!H{source_row}'
            ws[f"C{row}"].number_format = '#,##0.00'

        total_row = item_start + len(registros)
        aplicar_estilo_pagamento(total_row, "total")
        ws[f"A{total_row}"] = "Total"

        if len(registros):
            ws[f"C{total_row}"] = f"=SUM(C{item_start}:C{total_row - 1})"
        else:
            ws[f"C{total_row}"] = 0
        ws[f"C{total_row}"].number_format = '#,##0.00'

        # Uma linha em branco separa os dias.
        blank_row = total_row + 1
        aplicar_estilo_pagamento(blank_row, "blank")
        for c in range(1, 7):
            ws.cell(blank_row, c).value = None

        current_row = blank_row + 1

    # Os impostos não têm bloco separado: entram nas mesmas tabelas
    # pequenas de CONTAS À PAGAR, na data correspondente, identificados
    # por [IMPOSTO] no campo Nome.

    # Resumo de CONTAS À PAGAR: fornecedores + impostos, usando a data
    # original (sem D+). Assim o resumo principal apresenta uma única saída.
    for c_idx in range(2, len(dias) + 2):
        col = get_column_letter(c_idx)
        ws[f"{col}13"] = (
            f'=SUMIFS(Detalhamento!$H:$H,Detalhamento!$A:$A,{col}$4,'
            f'Detalhamento!$D:$D,"*FORNECED*")+'
            f'SUMIFS(Detalhamento!$H:$H,Detalhamento!$A:$A,{col}$4,'
            f'Detalhamento!$D:$D,"*IMPOSTOS*")'
        )
        ws[f"{col}13"].number_format = '#,##0.00'

    # A linha de IMPOSTOS deixa de existir no resumo principal. O detalhe
    # continua sendo gerado nas tabelas pequenas mais abaixo.
    ws["A14"] = None
    for c_idx in range(2, 7):
        ws.cell(14, c_idx).value = None

    # Saldo = saldo anterior + recebimentos - contas a pagar.
    for c_idx in range(2, len(dias) + 2):
        col = get_column_letter(c_idx)
        if c_idx == 2:
            ws[f"{col}15"] = "=B6+B8-B13"
        else:
            prev = get_column_letter(c_idx - 1)
            ws[f"{col}15"] = f"={prev}15+{col}8-{col}13"
        ws[f"{col}15"].number_format = '#,##0.00'

    # ------------------------------------------------------------
    # Configurações mantidas no modelo.
    # ------------------------------------------------------------
    if "Config_Carteiras" in wb.sheetnames:
        cws = wb["Config_Carteiras"]
        # Atualiza com a configuração usada pelo aplicativo.
        if cws.max_row > 1:
            cws.delete_rows(2, cws.max_row - 1)

        for _, r in config.iterrows():
            carteira = str(r.get("Carteira", "")).strip()
            if carteira:
                try:
                    dmais = int(r.get("D+", 1))
                except Exception:
                    dmais = 1
                cws.append([carteira, dmais, ""])
        cws["E1"] = ("Edite o D+ aqui (ou as datas em Dias_Sem_Compensacao): "
                     "Detalhamento e Planilha1 recalculam sozinhos. "
                     "Carteira sem cadastro = D+1.")

    if "Dias_Sem_Compensacao" in wb.sheetnames:
        fws = wb["Dias_Sem_Compensacao"]
        if fws.max_row > 1:
            fws.delete_rows(2, fws.max_row - 1)

        for d in sorted(sem_comp):
            fws.append([d, ""])
            fws.cell(fws.max_row, 1).number_format = "dd/mm/yyyy"

    # ------------------------------------------------------------
    # Salvar
    # ------------------------------------------------------------
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()

# ============================================================
# Interface
# ============================================================

st.title("💰 Previsão Financeira")
st.caption("Modelo baseado no layout da planilha de previsão utilizada atualmente.")

with st.sidebar:
    st.header("⚙️ Ferramentas")
    st.caption("Tudo o que você precisa para montar a previsão está aqui.")

    st.subheader("📥 Entrada de dados")
    uploaded = st.file_uploader("Carregar CSV bruto", type=["csv"], help="Selecione o arquivo CSV exportado do Tecnicon.")

    st.divider()
    st.subheader("📅 Previsão")
    saldo_inicial = st.number_input("Saldo inicial", min_value=0.0, value=50000.0, step=100.0, format="%.2f")
    data_inicio = st.date_input(
        "Data inicial",
        value=date.today() - timedelta(days=date.today().weekday()),
        format="DD/MM/YYYY",
    )
    data_fim = st.date_input(
        "Data final",
        value=data_inicio + timedelta(days=4),
        format="DD/MM/YYYY",
    )
    if data_fim < data_inicio:
        st.error("A data final não pode ser anterior à data inicial.")
        st.stop()
    if (data_fim - data_inicio).days > 4:
        st.error("A previsão pode ter no máximo 5 dias para manter o layout do Excel.")
        st.stop()

    with st.expander("💳 Carteiras / D+", expanded=False):
        cfg_default = pd.DataFrame({"Carteira": ["100", "33", "102", "74"], "D+": [1, 1, 0, 1]})
        config = st.data_editor(
            cfg_default,
            num_rows="dynamic",
            use_container_width=True,
            key="config",
            column_config={
                "Carteira": st.column_config.TextColumn("Carteira"),
                "D+": st.column_config.NumberColumn("D+", min_value=0, step=1),
            },
        )

    with st.expander("📅 Feriados", expanded=False):
        fer_default = pd.DataFrame({
            "Data": pd.Series(dtype="datetime64[ns]"),
            "Descrição": pd.Series(dtype="string"),
        })
        feriados = st.data_editor(
            fer_default,
            num_rows="dynamic",
            use_container_width=True,
            key="feriados",
            column_config={
                "Data": st.column_config.DateColumn("Data", format="DD/MM/YYYY"),
                "Descrição": st.column_config.TextColumn("Descrição"),
            },
        )

    # Espaço reservado para os downloads, preenchido depois do processamento.
    st.divider()
    st.subheader("📤 Exportação")
    download_area = st.empty()

if uploaded is None:
    st.info("Carregue o cubo bruto para gerar a previsão. Sem arquivo, o sistema mostra somente a estrutura.")
    st.stop()

try:
    df = ler_cubo_bruto(uploaded)
except Exception as e:
    st.error(f"Erro ao ler o CSV: {e}")
    st.stop()

if df.empty:
    st.error("Nenhum lançamento foi encontrado no CSV.")
    st.stop()

# Configuração de D+
map_d = {}
for _, r in config.iterrows():
    carteira = str(r.get("Carteira", "")).strip()
    if carteira:
        try:
            map_d[carteira] = int(r.get("D+", 1))
        except Exception:
            map_d[carteira] = 1

sem_comp = set()
if "Data" in feriados.columns:
    for x in feriados["Data"].dropna():
        try:
            sem_comp.add(x.date() if hasattr(x, "date") else x)
        except Exception:
            pass

# D+ não cadastrado usa D+1 como padrão, sem alerta na interface.
df["D+"] = df["Código"].astype(str).str.strip().map(map_d).fillna(1).astype(int)

df["Data Entrada"] = df.apply(
    lambda r: proxima_compensacao(r["Data"], int(r["D+"]), sem_comp)
    if eh_cliente(r["Tipo"]) else pd.NaT,
    axis=1,
)

df["Data Saída"] = df["Data"]

receber = df[df["Tipo"].map(eh_cliente)].copy()
pagar = df[df["Tipo"].map(eh_fornecedor)].copy()

# Período exibido: exatamente o intervalo escolhido.
segunda = data_inicio
dias = [data_inicio + timedelta(days=i) for i in range((data_fim - data_inicio).days + 1)]
nomes_semana = ["SEGUNDA", "TERÇA", "QUARTA", "QUINTA", "SEXTA", "SÁBADO", "DOMINGO"]
nomes = [nomes_semana[d.weekday()] for d in dias]

# ============================================================
# Painel no mesmo layout visual da planilha Excel
# ============================================================
from html import escape

DIAS_SEMANA_EXT = ["Segunda", "Terça", "Quarta", "Quinta",
                    "Sexta", "Sábado", "Domingo"]

def fmt_dia_semana(d):
    """Ex.: '29/09 - Terça' — usado em toda tabela que exibe datas."""
    d = pd.Timestamp(d)
    return f"{d.strftime('%d/%m')} - {DIAS_SEMANA_EXT[d.weekday()]}"

def _fmt_num(v):
    return f"{float(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

def _fmt_date(v):
    return pd.Timestamp(v).strftime("%d/%m/%Y")

def _recebimento_criterio(criterio, dia):
    return float(receber.loc[
        (receber["Carteira"] == criterio) & (receber["Data Entrada"] == dia),
        "Valor Cliente"
    ].sum())

recv_by_day = [float(receber.loc[receber["Data Entrada"] == d, "Valor Cliente"].sum()) for d in dias]
pay_by_day = [
    float(pagar.loc[pagar["Data Saída"] == d, "Valor Fornecedor"].sum())
    for d in dias
]

# Impostos são uma categoria própria do Cubo.
impostos = df[
    df["Tipo"].map(eh_imposto)
].copy()
impostos["Duplicata"] = impostos["Duplicata"].fillna("").astype(str).str.strip()
impostos["Fornecedor"] = impostos["Fornecedor"].fillna("").astype(str).str.strip()

impostos_por_dia = []
for d in dias:
    reg_imp = impostos[
        impostos["Data Saída"].eq(d)
        & (pd.to_numeric(impostos["Valor Fornecedor"], errors="coerce").fillna(0) != 0)
    ].copy()
    if not reg_imp.empty:
        reg_imp = (
            reg_imp.groupby(["Duplicata", "Fornecedor"], as_index=False)["Valor Fornecedor"]
            .sum()
            .sort_values(["Fornecedor", "Duplicata"], kind="stable")
        )
    impostos_por_dia.append(reg_imp)

impostos_by_day = [
    float(r["Valor Fornecedor"].sum()) if not r.empty else 0.0
    for r in impostos_por_dia
]

# CONTAS À PAGAR do resumo = fornecedores + impostos.
# Definimos antes do cálculo do saldo para evitar referência antecipada.
pay_total_by_day = [
    pay_by_day[i] + impostos_by_day[i] for i in range(len(dias))
]

saldos = []
saldo = float(saldo_inicial)
for i in range(len(dias)):
    saldo += recv_by_day[i] - pay_total_by_day[i]
    saldos.append(saldo)

fornecedores_por_dia = []
for d in dias:
    reg = pagar[
        pagar["Data Saída"].eq(d)
        & (pagar["Tipo"].map(eh_fornecedor) | pagar["Tipo"].map(eh_imposto))
        & (pd.to_numeric(pagar["Valor Fornecedor"], errors="coerce").fillna(0) != 0)
    ].copy()
    if not reg.empty:
        reg["Fornecedor"] = reg["Fornecedor"].fillna("").astype(str).str.strip()
        reg["Duplicata"] = reg["Duplicata"].fillna("").astype(str).str.strip()
        reg["Nome Exibicao"] = reg.apply(
            lambda r: ("[IMPOSTO] " if eh_imposto(r["Tipo"]) else "") + (r["Fornecedor"] or "(sem nome)"),
            axis=1,
        )
        reg = (
            reg.groupby(["Duplicata", "Nome Exibicao"], as_index=False)["Valor Fornecedor"]
            .sum()
            .sort_values(["Nome Exibicao", "Duplicata"], kind="stable")
        )
    fornecedores_por_dia.append(reg)

# Todas as carteiras de CLIENTES presentes no período são descobertas
# automaticamente. Uma carteira nova não precisa ser catalogada no código.
carteiras_recebimento = sorted(
    receber.loc[
        receber["Data Entrada"].isin(dias), "Carteira"
    ].dropna().astype(str).str.strip().replace("", pd.NA).dropna().unique().tolist()
)

rows_html = []
total_cols = len(dias) + 1
rows_html.append(f'<tr class="titulo"><td colspan="{total_cols}">PREVISÃO FINANCEIRA SEMANAL</td></tr>')
rows_html.append('<tr class="dias"><td></td>' + ''.join(f'<td>{fmt_dia_semana(dias[i])}</td>' for i in range(len(dias)) ) + '</tr>')
rows_html.append(f'<tr class="espaco"><td colspan="{total_cols}"></td></tr>')
rows_html.append(
    '<tr class="saldo-inicial"><td>SALDO INICIAL</td>'
    + f'<td class="num">{_fmt_num(saldo_inicial)}</td>'
    + ''.join('<td></td>' for _ in range(max(0, len(dias) - 1)))
    + '</tr>'
)
rows_html.append(f'<tr class="espaco"><td colspan="{total_cols}"></td></tr>')
rows_html.append('<tr class="secao"><td>CONTAS À RECEBER</td>' + ''.join(f'<td class="num">{_fmt_num(v)}</td>' for v in recv_by_day) + '</tr>')
for carteira in carteiras_recebimento:
    vals = [_recebimento_criterio(carteira, d) for d in dias]
    rows_html.append(
        f'<tr class="sub"><td>{escape(carteira)}</td>'
        + ''.join(f'<td class="num">{_fmt_num(v)}</td>' for v in vals)
        + '</tr>'
    )
rows_html.append(f'<tr class="espaco"><td colspan="{total_cols}"></td></tr>')
rows_html.append('<tr class="secao"><td>CONTAS À PAGAR</td>' + ''.join(f'<td class="num">{_fmt_num(v)}</td>' for v in pay_total_by_day) + '</tr>')
rows_html.append(f'<tr class="espaco"><td colspan="{total_cols}"></td></tr>')
rows_html.append('<tr class="saldo"><td>SALDO</td>' + ''.join(f'<td class="num">{_fmt_num(v)}</td>' for v in saldos) + '</tr>')
rows_html.append(f'<tr class="espaco"><td colspan="{total_cols}"></td></tr>')
rows_html.append(f'<tr class="secao"><td colspan="{total_cols}">CONTAS À PAGAR</td></tr>')

for idx, d in enumerate(dias):
    reg = fornecedores_por_dia[idx]
    rows_html.append(f'<tr class="data"><td>{fmt_dia_semana(d)}</td><td colspan="{max(1, total_cols-1)}"></td></tr>')
    # A tabela de detalhes usa sempre a mesma coluna de valor, independentemente
    # do dia. Assim os números não ficam pulando de uma coluna para outra.
    valor_colspan = max(1, total_cols - 2)
    rows_html.append(
        f'<tr class="cabecalho-mini"><td>Duplicata</td>'
        f'<td colspan="{valor_colspan}">Nome</td><td>Valor</td></tr>'
    )
    for _, r in reg.iterrows():
        rows_html.append(
            f'<tr class="fornecedor"><td>{escape(str(r["Duplicata"]))}</td>'
            f'<td colspan="{valor_colspan}">{escape(str(r["Nome Exibicao"]).strip() or "(sem nome)")}</td>'
            f'<td class="num">{_fmt_num(r["Valor Fornecedor"])}</td></tr>'
        )
    total = float(reg["Valor Fornecedor"].sum()) if not reg.empty else 0.0
    rows_html.append(
        f'<tr class="total"><td>Total</td><td colspan="{valor_colspan}"></td>'
        f'<td class="num">{_fmt_num(total)}</td></tr>'
    )

rows_html.append(f'<tr class="espaco"><td colspan="{total_cols}"></td></tr>')

html = f'''
<style>
.previsao-wrap {{ width:100%; overflow-x:auto; margin-top:4px; }}
.previsao {{ border-collapse:collapse; width:100%; min-width:820px; font-family:Calibri,Arial,sans-serif; font-size:14px; color:#111; background:#fff; }}
.previsao td {{ padding:4px 8px; height:24px; border:0; vertical-align:middle; }}
.previsao td:first-child {{ width:42%; text-align:left; }}
.previsao td:not(:first-child) {{ width:11.6%; text-align:right; }}
.previsao .titulo td {{ text-align:center; font-weight:700; font-size:16px; padding:9px 0 5px; }}
.previsao .dias td {{ text-align:center; font-weight:700; border-bottom:1px solid #222; }}
.previsao .dias span {{ font-weight:400; }}
.previsao .secao td {{ font-weight:700; border-bottom:1px solid #222; padding-top:7px; }}
.previsao .saldo-inicial td {{ font-weight:700; border-bottom:1px solid #222; padding:7px 8px; }}
.previsao .saldo-inicial .num {{ text-align:right; }}
.previsao .sub td {{ border:0; }}
.previsao .saldo td {{ font-weight:700; border-top:1px solid #222; border-bottom:1px solid #222; padding:6px 8px; }}
.previsao .data td {{ font-weight:700; padding-top:10px; border-bottom:1px solid #222; }}
.previsao .fornecedor td:first-child {{ white-space:nowrap; }}
 .previsao .fornecedor td:nth-child(2) {{ white-space:nowrap; overflow:hidden; text-overflow:ellipsis; }}
 .previsao .cabecalho-mini td {{ font-size:12px; font-weight:700; border-bottom:1px solid #aaa; }}
.previsao .total td {{ border-bottom:1px solid #222; font-weight:700; }}
.previsao .espaco td {{ height:8px; padding:0; }}
.previsao .num {{ font-variant-numeric:tabular-nums; }}
</style>
<div class="previsao-wrap"><table class="previsao">{''.join(rows_html)}</table></div>
'''
st.markdown(html, unsafe_allow_html=True)

st.divider()

st.subheader("🔎 Detalhamento do cálculo")
det = df[["Data", "Carteira", "Código", "Tipo", "D+", "Valor Cliente", "Data Entrada", "Valor Fornecedor", "Data Saída", "Duplicata", "Fornecedor"]].copy()
det["Classificação"] = det["Tipo"].map(lambda x: "Cliente" if eh_cliente(x) else ("A pagar" if eh_fornecedor(x) else ("Imposto" if eh_imposto(x) else "Não classificado")))
det.columns = ["Data bruto", "Carteira", "Código", "Tipo", "D+", "Cliente", "Entrada no banco", "A pagar", "Saída", "Duplicata", "Fornecedor", "Classificação"]
st.dataframe(det.sort_values(["Data bruto", "Carteira"]), use_container_width=True, hide_index=True)

# Downloads
excel_bytes = gerar_excel(df, saldo_inicial, data_inicio, config, sem_comp, data_fim=data_fim)
download_area.download_button(
    "📊 Baixar previsão em Excel",
    data=excel_bytes,
    file_name=f"Previsao_Financeira_{data_inicio.strftime('%Y%m%d')}_{data_fim.strftime('%Y%m%d')}.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    use_container_width=True,
)

csv_result = pd.DataFrame({
    "Data": [fmt_dia_semana(d) for d in dias],
    "A Receber": recv_by_day,
    "A Pagar": pay_total_by_day,
    "Saldo": saldos,
})
csv_bytes = csv_result.to_csv(index=False, sep=";", decimal=",", encoding="utf-8-sig").encode("utf-8-sig")
download_area.download_button(
    "📄 Baixar resultado CSV",
    data=csv_bytes,
    file_name=f"Previsao_Financeira_{data_inicio.strftime('%Y%m%d')}_{data_fim.strftime('%Y%m%d')}.csv",
    mime="text/csv",
    use_container_width=True,
)
